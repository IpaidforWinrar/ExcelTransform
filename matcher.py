"""
Befüllt "empty Auswertung Benutzergruppen ZB IT.xlsx" und ordnet jede Gruppe einem Bereich zu.

Ablauf:
1. Spalte A von "Auswertung Benutzergruppen_bearbeitet.xlsx" (Ergebnis von transform.py)
   kommt in Spalte B (Gruppe), Spalte B kommt in Spalte C (Name) der leeren Vorlage.
   Gelesen wird ab Zeile 2 (Zeile 1 ist die Überschrift), eingefügt ab Zeile 3.
2. Für jede Gruppe wird in "Benutzergruppen Zuordnung.xlsx" gesucht. Zeile 1 dort enthält die
   Bereiche (IT 1, IT 2, ..., IT-Leitung + Stab, übergreifend), darunter die Gruppen.
   Bei einem Treffer kommt die Spaltenüberschrift in Spalte A (betrifft).
3. Ohne Treffer wird der Gruppenname geprüft: IT_1... -> IT 1, IT_6... -> IT 6,
   IT_0... -> IT-Leitung + Stab.
4. Klappt auch das nicht, werden alle offenen Gruppen im Terminal zuerst als Liste
   ausgegeben und danach einzeln abgefragt.

Neue Zuordnungen aus Schritt 3 und 4 werden in "Benutzergruppen Zuordnung.xlsx" übernommen,
damit sie beim nächsten Mal direkt gefunden werden. Vorher wird davon eine Sicherung angelegt.

Das Ergebnis wird als "Auswertung Benutzergruppen ZB IT.xlsx" gespeichert, die leere Vorlage
bleibt unverändert.

Aufruf:
    python matcher.py

Die Exceldateien werden im Ordner dieses Skripts gesucht, danach im aktuellen Ordner.
"""

import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ORDNER = Path(__file__).resolve().parent

QUELLE_DATEI = "Auswertung Benutzergruppen_bearbeitet.xlsx"
VORLAGE_DATEI = "empty Auswertung Benutzergruppen ZB IT.xlsx"
ZUORDNUNG_DATEI = "Benutzergruppen Zuordnung.xlsx"
AUSGABE_NAME = "Auswertung Benutzergruppen ZB IT.xlsx"

QUELLE_START_ROW = 2  # Zeile 1 ist die Spaltenüberschrift
VORLAGE_START_ROW = 3

COL_BETRIFFT = 1  # Spalte A
COL_GRUPPE = 2  # Spalte B
COL_NAME = 3  # Spalte C

# IT_0 gehört zu "IT-Leitung + Stab", IT_1 bis IT_9 zu "IT 1" bis "IT 9"
IT_MUSTER = re.compile(r"(?<![a-z0-9])it_(\d)", re.IGNORECASE)
LEITUNG = "IT-Leitung + Stab"


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def gruppe_key(text: str) -> str:
    return clean(text).casefold()


def bereich_key(text: str) -> str:
    # "IT1", "it 1" und "IT  1" werden alle als "it 1" erkannt
    return re.sub(r"^it\s*(\d)$", r"it \1", clean(text).casefold())


def find_file(name: str) -> Path:
    for ordner in (ORDNER, Path.cwd()):
        path = ordner / name
        if path.exists():
            return path
    sys.exit(f"Datei nicht gefunden: {name} (gesucht in {ORDNER} und {Path.cwd()})")


class Zuordnung:
    """Liest "Benutzergruppen Zuordnung.xlsx" und kann neue Gruppen in die passende Spalte schreiben."""

    def __init__(self, path: Path):
        self.path = path
        self.wb = load_workbook(path)
        self.ws = self.wb.active
        self.geaendert = False

        # Bereich -> Spaltennummer, in der Reihenfolge der Tabelle
        self.spalten = {}
        for cell in self.ws[1]:
            header = clean(cell.value)
            if header:
                self.spalten[header] = cell.column

        # Gruppe -> Liste der Bereiche, in denen sie steht
        self.gruppen = {}
        for bereich, col in self.spalten.items():
            for row in range(2, self.ws.max_row + 1):
                gruppe = clean(self.ws.cell(row=row, column=col).value)
                if gruppe:
                    self.gruppen.setdefault(gruppe_key(gruppe), [])
                    if bereich not in self.gruppen[gruppe_key(gruppe)]:
                        self.gruppen[gruppe_key(gruppe)].append(bereich)

    @property
    def bereiche(self):
        return list(self.spalten)

    def finde_bereich(self, name: str):
        """Sucht einen Bereich anhand des Namens, z. B. "IT 1" oder "IT-Leitung + Stab"."""
        for bereich in self.spalten:
            if bereich_key(bereich) == bereich_key(name):
                return bereich
        return None

    def lookup(self, gruppe: str):
        return self.gruppen.get(gruppe_key(gruppe), [])

    def hinzufuegen(self, gruppe: str, bereich: str) -> None:
        col = self.spalten[bereich]
        row = self.ws.max_row
        while row > 1 and clean(self.ws.cell(row=row, column=col).value) == "":
            row -= 1
        self.ws.cell(row=row + 1, column=col, value=gruppe)
        self.gruppen.setdefault(gruppe_key(gruppe), []).append(bereich)
        self.geaendert = True

    def speichern(self) -> None:
        if not self.geaendert:
            return
        backup = self.path.with_name(f"{self.path.stem}_backup{self.path.suffix}")
        shutil.copy2(self.path, backup)
        self.wb.save(self.path)
        print(f"Neue Zuordnungen übernommen in: {self.path} (Sicherung: {backup.name})")


def bereich_aus_namen(gruppe: str, zuordnung: Zuordnung):
    treffer = IT_MUSTER.search(gruppe)
    if not treffer:
        return None
    ziffer = treffer.group(1)
    name = LEITUNG if ziffer == "0" else f"IT {ziffer}"
    return zuordnung.finde_bereich(name)


def frage_bereich(gruppe: str, name: str, row: int, zuordnung: Zuordnung, hinweis: str = ""):
    bereiche = zuordnung.bereiche
    print()
    print(f"Keine Zuordnung für Zeile {row}: Gruppe '{gruppe}'" + (f", Name '{name}'" if name else ""))
    if hinweis:
        print(hinweis)
    for idx, bereich in enumerate(bereiche, start=1):
        print(f"  {idx}) {bereich}")
    while True:
        try:
            antwort = input(f"Bereich wählen (1-{len(bereiche)}, Enter = überspringen): ").strip()
        except EOFError:
            return None
        if antwort == "":
            return None
        if antwort.isdigit() and 1 <= int(antwort) <= len(bereiche):
            return bereiche[int(antwort) - 1]
        print("Ungültige Eingabe.")


def read_quelle(path: Path):
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = []
    for row in ws.iter_rows(min_row=QUELLE_START_ROW, max_col=2, values_only=True):
        gruppe = row[0] if len(row) > 0 else None
        name = row[1] if len(row) > 1 else None
        if clean(gruppe) == "" and clean(name) == "":
            continue
        rows.append((gruppe, name))
    return rows


def main() -> None:
    quelle_datei = find_file(QUELLE_DATEI)
    vorlage_datei = find_file(VORLAGE_DATEI)
    zuordnung_datei = find_file(ZUORDNUNG_DATEI)
    ausgabe_datei = vorlage_datei.with_name(AUSGABE_NAME)

    zuordnung = Zuordnung(zuordnung_datei)
    if not zuordnung.bereiche:
        sys.exit(f"Keine Bereiche in Zeile 1 von {zuordnung_datei.name} gefunden.")

    # Schritt 1: Gruppen und Namen in die Vorlage übertragen
    wb = load_workbook(vorlage_datei)
    ws = wb.active
    quelle = read_quelle(quelle_datei)
    for offset, (gruppe, name) in enumerate(quelle):
        row = VORLAGE_START_ROW + offset
        ws.cell(row=row, column=COL_GRUPPE, value=gruppe)
        ws.cell(row=row, column=COL_NAME, value=name)

    # Schritt 2 und 3: Bereich automatisch in Spalte A eintragen
    zaehler = {"Treffer": 0, "aus IT_X": 0, "manuell": 0, "ohne Zuordnung": 0}
    offen = []  # (row, gruppe, name), die nicht automatisch zugeordnet werden konnten
    for row in range(VORLAGE_START_ROW, VORLAGE_START_ROW + len(quelle)):
        gruppe = clean(ws.cell(row=row, column=COL_GRUPPE).value)
        name = clean(ws.cell(row=row, column=COL_NAME).value)
        if not gruppe:
            continue

        gefunden = zuordnung.lookup(gruppe)
        if len(gefunden) == 1:
            ws.cell(row=row, column=COL_BETRIFFT, value=gefunden[0])
            zaehler["Treffer"] += 1
            continue

        bereich = bereich_aus_namen(gruppe, zuordnung) if not gefunden else None
        if bereich:
            ws.cell(row=row, column=COL_BETRIFFT, value=bereich)
            zuordnung.hinzufuegen(gruppe, bereich)
            zaehler["aus IT_X"] += 1
            continue

        offen.append((row, gruppe, name))

    # Schritt 4: nicht zuordenbare Gruppen erst auflisten, dann einzeln abfragen
    if offen:
        print(f"{len(offen)} Gruppen konnten nicht automatisch zugeordnet werden:")
        for row, gruppe, name in offen:
            gefunden = zuordnung.lookup(gruppe)
            zusatz = f"  [steht in mehreren Bereichen: {', '.join(gefunden)}]" if gefunden else ""
            print(f"  Zeile {row}: {gruppe}" + (f" ({name})" if name else "") + zusatz)

    for row, gruppe, name in offen:
        gefunden = zuordnung.lookup(gruppe)
        # Gleiche Gruppe wurde weiter oben schon manuell zugeordnet
        if len(gefunden) == 1:
            ws.cell(row=row, column=COL_BETRIFFT, value=gefunden[0])
            zaehler["manuell"] += 1
            continue

        hinweis = f"Steht in mehreren Bereichen: {', '.join(gefunden)}" if gefunden else ""
        bereich = frage_bereich(gruppe, name, row, zuordnung, hinweis)
        if not bereich:
            zaehler["ohne Zuordnung"] += 1
            continue

        ws.cell(row=row, column=COL_BETRIFFT, value=bereich)
        zaehler["manuell"] += 1
        # Gruppen, die schon in mehreren Bereichen stehen, nicht noch einmal eintragen
        if not gefunden:
            zuordnung.hinzufuegen(gruppe, bereich)

    wb.save(ausgabe_datei)
    zuordnung.speichern()

    print()
    print(f"{len(quelle)} Zeilen übertragen.")
    for art, anzahl in zaehler.items():
        print(f"  {art}: {anzahl}")
    print(f"Fertig: {ausgabe_datei}")


if __name__ == "__main__":
    main()
