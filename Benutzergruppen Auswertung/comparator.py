"""
Vergleicht "filled Auswertung Benutzergruppen ZB IT.xlsx" (dieses Jahr) mit
"toCompareLastYear.xlsx" (letztes Jahr) und listet Neuzugänge und Abgänge.

Beide Tabellen haben denselben Aufbau: A = betrifft, B = Gruppe, C = Name, Werte ab Zeile 3.

Ein Neuzugang ist eine Kombination aus Gruppe und Name, die dieses Jahr vorkommt,
letztes Jahr aber nicht. Ein Abgang ist umgekehrt eine Kombination, die letztes Jahr
vorkam, dieses Jahr aber nicht mehr. Groß-/Kleinschreibung und überzählige Leerzeichen
werden ignoriert. Ändert sich nur "betrifft", zählt das weder als Neuzugang noch als Abgang.

Ergebnis ist "Veränderungen.xlsx" mit den Reitern "Neuzugänge" und "Abgänge",
jeweils mit den Spalten betrifft, Gruppe, Name. Bei Abgängen steht in "betrifft"
der Bereich vom letzten Jahr.

Aufruf:
    python comparator.py

"toCompareLastYear.xlsx" liegt in "FAVIS DB Benutzergruppen/Eingabe Benutzergruppen",
die Datei von diesem Jahr und das Ergebnis in "FAVIS DB Benutzergruppen/Ausgabe Benutzergruppen".
"""

import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ORDNER = Path(__file__).resolve().parent

EINGABE_ORDNER = ORDNER.parent / "FAVIS DB Benutzergruppen" / "Eingabe Benutzergruppen"
AUSGABE_ORDNER = ORDNER.parent / "FAVIS DB Benutzergruppen" / "Ausgabe Benutzergruppen"

AKTUELL_DATEI = AUSGABE_ORDNER / "filled Auswertung Benutzergruppen ZB IT.xlsx"
VORJAHR_DATEI = EINGABE_ORDNER / "toCompareLastYear.xlsx"
AUSGABE_DATEI = AUSGABE_ORDNER / "Veränderungen.xlsx"

START_ROW = 3
HEADERS = ["betrifft", "Gruppe", "Name"]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def find_file(path: Path) -> Path:
    if not path.exists():
        sys.exit(f"Datei nicht gefunden: {path}")
    return path


def read_rows(path: Path):
    """Liefert (betrifft, gruppe, name) für jede Zeile ab START_ROW mit Gruppe oder Name."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = []
    for row in ws.iter_rows(min_row=START_ROW, max_col=3, values_only=True):
        betrifft, gruppe, name = (clean(row[i]) if i < len(row) else "" for i in range(3))
        if gruppe or name:
            rows.append((betrifft, gruppe, name))
    return rows


def key(gruppe: str, name: str):
    return (gruppe.casefold(), name.casefold())


def nur_in(rows, andere):
    """Zeilen aus rows, deren Gruppe und Name in andere nicht vorkommen, ohne Wiederholungen."""
    andere_keys = {key(gruppe, name) for _, gruppe, name in andere}
    result = []
    gesehen = set()
    for betrifft, gruppe, name in rows:
        k = key(gruppe, name)
        if k in andere_keys or k in gesehen:
            continue
        gesehen.add(k)
        result.append((betrifft, gruppe, name))
    return result


def write_sheet(ws, title: str, rows) -> None:
    ws.title = title
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)

    for col_idx, header in enumerate(HEADERS, start=1):
        width = max([len(header)] + [len(r[col_idx - 1]) for r in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    aktuell_datei = find_file(AKTUELL_DATEI)
    vorjahr_datei = find_file(VORJAHR_DATEI)
    ausgabe_datei = AUSGABE_DATEI

    aktuell = read_rows(aktuell_datei)
    vorjahr = read_rows(vorjahr_datei)

    neuzugaenge = nur_in(aktuell, vorjahr)
    abgaenge = nur_in(vorjahr, aktuell)

    wb = Workbook()
    write_sheet(wb.active, "Neuzugänge", neuzugaenge)
    write_sheet(wb.create_sheet(), "Abgänge", abgaenge)
    wb.save(ausgabe_datei)

    print(f"Dieses Jahr: {len(aktuell)} Zeilen, letztes Jahr: {len(vorjahr)} Zeilen")
    print(f"Neuzugänge: {len(neuzugaenge)}")
    print(f"Abgänge: {len(abgaenge)}")
    print(f"Fertig: {ausgabe_datei}")


if __name__ == "__main__":
    main()
