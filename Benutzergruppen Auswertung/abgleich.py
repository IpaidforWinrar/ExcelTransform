"""
Gleicht test_Zuordnung.xlsx (Ergebnis von zuordnung.py) mit matching.xlsx ab.

test_Zuordnung.xlsx: Zeile 1 enthält die Bereiche (IT 1 bis IT 6 usw.),
darunter stehen die Begriffe, die dem jeweiligen Bereich zugeordnet sind.

matching.xlsx: ab Zeile 2 steht in Spalte A der Begriff und in Spalte F der Bereich.

Verglichen werden die Paare (Bereich, Begriff). Groß-/Kleinschreibung und
überzählige Leerzeichen werden dabei ignoriert. Ausgabe ist Abgleich.xlsx mit
einer Zeile pro Paar und dem Status:
    - "in beiden"
    - "nur in test_Zuordnung"
    - "nur in matching"
Die Spalte "In der anderen Liste bei" zeigt, welchem Bereich der Begriff in der
anderen Liste zugeordnet ist, falls das Paar dort fehlt.

Zusätzlich entsteht Abweichungen.xlsx mit einer Zeile pro auffälligem Begriff:
    - "Widerspruch": der Begriff ist in beiden Listen unterschiedlichen Bereichen zugeordnet
    - "fehlt in matching" / "fehlt in test_Zuordnung": der Begriff steht nur in einer Liste

Aufruf:
    python abgleich.py

Die Exceldateien werden im Ordner dieses Skripts gesucht, danach im aktuellen Ordner.
Die Ergebnisse landen im selben Ordner wie matching.xlsx.
"""

import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ORDNER = Path(__file__).resolve().parent

BEREICHE = [
    "IT 1",
    "IT 2",
    "IT 3",
    "IT 4",
    "IT 5",
    "IT 6",
    "IT-Leitung + Stab",
    "übergreifend",
]

MATCHING_START_ROW = 2
MATCHING_COL_BEGRIFF = 1  # Spalte A
MATCHING_COL_BEREICH = 6  # Spalte F

STATUS_BEIDE = "in beiden"
STATUS_NUR_ZUORDNUNG = "nur in test_Zuordnung"
STATUS_NUR_MATCHING = "nur in matching"

STATUS_FARBEN = {
    STATUS_BEIDE: "C6EFCE",
    STATUS_NUR_ZUORDNUNG: "FFEB9C",
    STATUS_NUR_MATCHING: "FFC7CE",
}


ART_WIDERSPRUCH = "Widerspruch"
ART_FEHLT_MATCHING = "fehlt in matching"
ART_FEHLT_ZUORDNUNG = "fehlt in test_Zuordnung"

ART_FARBEN = {
    ART_WIDERSPRUCH: "FFC7CE",
    ART_FEHLT_MATCHING: "FFEB9C",
    ART_FEHLT_ZUORDNUNG: "FFEB9C",
}


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_key(text: str) -> str:
    # "IT1", "it 1" und "IT  1" werden alle als "it 1" erkannt
    return re.sub(r"^it\s*(\d)$", r"it \1", text.lower())


BEREICH_KEYS = {normalize_key(b): b for b in BEREICHE}


def find_file(name: str) -> Path:
    for ordner in (ORDNER, Path.cwd()):
        path = ordner / name
        if path.exists():
            return path
    sys.exit(f"Datei nicht gefunden: {name} (gesucht in {ORDNER} und {Path.cwd()})")


def bereich_name(text: str) -> str:
    return BEREICH_KEYS.get(normalize_key(text), text)


def begriff_key(text: str) -> str:
    return text.casefold()


def read_zuordnung(path: Path):
    """Liefert {(bereich, begriff_key): begriff} aus test_Zuordnung.xlsx."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    pairs = {}
    if not rows:
        return pairs
    headers = [clean(h) for h in rows[0]]
    for row in rows[1:]:
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            begriff = clean(row[col_idx] if col_idx < len(row) else None)
            if begriff:
                pairs[(bereich_name(header), begriff_key(begriff))] = begriff
    return pairs


def read_matching(path: Path):
    """Liefert {(bereich, begriff_key): begriff} aus matching.xlsx und übersprungene Zeilen."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    pairs = {}
    skipped = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=MATCHING_START_ROW, max_col=MATCHING_COL_BEREICH, values_only=True),
        start=MATCHING_START_ROW,
    ):
        begriff = clean(row[MATCHING_COL_BEGRIFF - 1] if len(row) >= MATCHING_COL_BEGRIFF else None)
        bereich = clean(row[MATCHING_COL_BEREICH - 1] if len(row) >= MATCHING_COL_BEREICH else None)
        if not begriff and not bereich:
            continue
        if not begriff or not bereich:
            skipped.append(row_idx)
            continue
        pairs[(bereich_name(bereich), begriff_key(begriff))] = begriff
    return pairs, skipped


def bereiche_pro_begriff(pairs):
    result = defaultdict(set)
    for bereich, key in pairs:
        result[key].add(bereich)
    return result


def sort_key_bereich(bereich: str):
    return (BEREICHE.index(bereich) if bereich in BEREICHE else len(BEREICHE), bereich)


def compare(zuordnung, matching):
    bereiche_zuordnung = bereiche_pro_begriff(zuordnung)
    bereiche_matching = bereiche_pro_begriff(matching)

    rows = []
    for pair in set(zuordnung) | set(matching):
        bereich, key = pair
        if pair in zuordnung and pair in matching:
            status, andere = STATUS_BEIDE, set()
        elif pair in zuordnung:
            status, andere = STATUS_NUR_ZUORDNUNG, bereiche_matching.get(key, set())
        else:
            status, andere = STATUS_NUR_MATCHING, bereiche_zuordnung.get(key, set())
        begriff = zuordnung.get(pair) or matching.get(pair)
        andere_text = ", ".join(sorted(andere, key=sort_key_bereich))
        rows.append((begriff, bereich, status, andere_text))

    status_order = [STATUS_NUR_ZUORDNUNG, STATUS_NUR_MATCHING, STATUS_BEIDE]
    rows.sort(key=lambda r: (status_order.index(r[2]), sort_key_bereich(r[1]), r[0].casefold()))
    return rows


def find_abweichungen(zuordnung, matching):
    """Eine Zeile pro Begriff, dessen Bereiche in beiden Listen nicht übereinstimmen."""
    bereiche_zuordnung = bereiche_pro_begriff(zuordnung)
    bereiche_matching = bereiche_pro_begriff(matching)
    namen = {key: begriff for (_, key), begriff in {**matching, **zuordnung}.items()}

    rows = []
    for key in set(bereiche_zuordnung) | set(bereiche_matching):
        in_zuordnung = bereiche_zuordnung.get(key, set())
        in_matching = bereiche_matching.get(key, set())
        if in_zuordnung == in_matching:
            continue
        if not in_matching:
            art = ART_FEHLT_MATCHING
        elif not in_zuordnung:
            art = ART_FEHLT_ZUORDNUNG
        else:
            art = ART_WIDERSPRUCH
        rows.append((
            namen[key],
            ", ".join(sorted(in_zuordnung, key=sort_key_bereich)),
            ", ".join(sorted(in_matching, key=sort_key_bereich)),
            art,
        ))

    art_order = [ART_WIDERSPRUCH, ART_FEHLT_MATCHING, ART_FEHLT_ZUORDNUNG]
    rows.sort(key=lambda r: (art_order.index(r[3]), r[0].casefold()))
    return rows


def write_abweichungen(rows, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Abweichungen"

    headers = ["Begriff", "Bereich in test_Zuordnung", "Bereich in matching", "Art"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)
        ws.cell(row=ws.max_row, column=4).fill = PatternFill("solid", start_color=ART_FARBEN[row[3]])

    for col_idx, header in enumerate(headers, start=1):
        width = max([len(header)] + [len(str(r[col_idx - 1])) for r in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_output(rows, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Abgleich"

    headers = ["Begriff", "Bereich", "Status", "In der anderen Liste bei"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)
        ws.cell(row=ws.max_row, column=3).fill = PatternFill(
            "solid", start_color=STATUS_FARBEN[row[2]]
        )

    for col_idx, header in enumerate(headers, start=1):
        width = max([len(header)] + [len(str(r[col_idx - 1])) for r in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    zuordnung_datei = find_file("test_Zuordnung.xlsx")
    matching_datei = find_file("matching.xlsx")
    abgleich_datei = matching_datei.with_name("Abgleich.xlsx")
    abweichungen_datei = matching_datei.with_name("Abweichungen.xlsx")

    zuordnung = read_zuordnung(zuordnung_datei)
    matching, skipped = read_matching(matching_datei)
    print(f"Gelesen: {len(zuordnung)} Paare aus {zuordnung_datei.name}, {len(matching)} Paare aus {matching_datei.name}")

    rows = compare(zuordnung, matching)
    write_output(rows, abgleich_datei)

    abweichungen = find_abweichungen(zuordnung, matching)
    write_abweichungen(abweichungen, abweichungen_datei)

    for status in (STATUS_BEIDE, STATUS_NUR_ZUORDNUNG, STATUS_NUR_MATCHING):
        print(f"{status}: {sum(1 for r in rows if r[2] == status)}")
    for art in (ART_WIDERSPRUCH, ART_FEHLT_MATCHING, ART_FEHLT_ZUORDNUNG):
        print(f"{art}: {sum(1 for r in abweichungen if r[3] == art)}")
    if skipped:
        print(f"Warnung: Zeilen in matching.xlsx mit nur einem Wert in A oder F übersprungen: {', '.join(map(str, skipped))}")
    print(f"Fertig: {abgleich_datei}")
    print(f"Fertig: {abweichungen_datei}")


if __name__ == "__main__":
    main()
