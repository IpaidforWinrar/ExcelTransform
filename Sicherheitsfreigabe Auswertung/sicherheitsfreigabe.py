"""
Ordnet die Zeilen von "Auswertung Sicherheitsfreigabe.xlsx" den IT-Bereichen zu und
erstellt pro Bereich eine eigene Liste.

Eingabe:
- "Auswertung Sicherheitsfreigabe.xlsx": Spalten A bis H, Zeile 1 = Überschrift, Werte ab Zeile 2.
  In Spalte H stehen mit "," getrennte Gruppen.
- "Benutzergruppen Zuordnung.xlsx": Zeile 1 = Bereiche (IT 1 bis IT 7, IT-Leitung + Stab,
  übergreifend), darunter die Gruppen. Wird nur gelesen, nie verändert.
- "empty Auswertung Sicherheitsfreigabe.xlsx": Vorlage, Überschriften in Zeile 2, Werte ab Zeile 3,
  zusätzliche Spalte I.

Zuordnung pro Zeile:
- Jede Gruppe aus Spalte H wird in "Benutzergruppen Zuordnung.xlsx" gesucht. Ohne Treffer wird der
  Name geprüft: IT_1... -> IT 1 bis IT_7... -> IT 7, IT_0... -> IT-Leitung + Stab.
  Gruppen ohne jeden Treffer werden ignoriert.
- Die Zeile bekommt alle Bereiche, die dabei gefunden werden. Sind es mehrere, stehen sie in
  Spalte J "Doppelzuordnung".

Ausgabe: 10 Dateien "Auswertung Sicherheitsfreigabe und Gruppenzuordnung ZB IT - X.xlsx" mit
X = IT 1 bis IT 7, IT-L_Stab, IT-übergreifend und IT - gesamt. Jede Bereichsliste enthält alle Zeilen,
die diesem Bereich zugeordnet sind, auch mehrfach zugeordnete. Die Gesamtliste enthält alle Zeilen
und zusätzlich Spalte K "Zuordnung".

Aufruf:
    python sicherheitsfreigabe.py

Die Exceldateien werden im Ordner dieses Skripts gesucht, danach im Ordner "Benutzergruppen
Auswertung" daneben, im übergeordneten Ordner und im aktuellen Ordner. Die Ergebnisse landen im
Unterordner "Ausgabe Auswertung", der bei Bedarf angelegt wird. Dort liegt auch die Logdatei
"Auswertung Sicherheitsfreigabe.log" mit allem, was im Terminal ausgegeben wurde.
"""

import re
import sys
import traceback
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ORDNER = Path(__file__).resolve().parent
# "Benutzergruppen Zuordnung.xlsx" wird von matcher.py im Nachbarordner gepflegt
BENUTZERGRUPPEN_ORDNER = ORDNER.parent / "Benutzergruppen Auswertung"
AUSGABE_ORDNER = ORDNER / "Ausgabe Auswertung"

QUELLE_DATEI = "Auswertung Sicherheitsfreigabe.xlsx"
ZUORDNUNG_DATEI = "Benutzergruppen Zuordnung.xlsx"
VORLAGE_DATEI = "empty Auswertung Sicherheitsfreigabe.xlsx"
AUSGABE_MUSTER = "Auswertung Sicherheitsfreigabe und Gruppenzuordnung ZB IT - {}.xlsx"
LOG_DATEI = AUSGABE_ORDNER / "Auswertung Sicherheitsfreigabe.log"

QUELLE_START_ROW = 2
QUELLE_SPALTEN = 8  # A bis H
COL_GRUPPEN = 8  # Spalte H

VORLAGE_HEADER_ROW = 2
VORLAGE_START_ROW = 3
COL_DOPPELT = 10  # Spalte J
COL_ZUORDNUNG = 11  # Spalte K, nur in der Gesamtliste

# Bereich -> Kürzel im Dateinamen
BEREICHE = {
    "IT 1": "IT 1",
    "IT 2": "IT 2",
    "IT 3": "IT 3",
    "IT 4": "IT 4",
    "IT 5": "IT 5",
    "IT 6": "IT 6",
    "IT 7": "IT 7",
    "IT-Leitung + Stab": "IT-L_Stab",
    "übergreifend": "IT-übergreifend",
}
GESAMT = "IT - gesamt"
LEITUNG = "IT-Leitung + Stab"

# IT_0 gehört zu "IT-Leitung + Stab", IT_1 bis IT_7 zu "IT 1" bis "IT 7"
IT_MUSTER = re.compile(r"(?<![a-z0-9])it_(\d)", re.IGNORECASE)


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def gruppe_key(text: str) -> str:
    return clean(text).casefold()


def bereich_key(text: str) -> str:
    # "IT 1", "IT1" und "it-1" sowie "IT-Leitung + Stab" und "IT Leitung + Stab" gelten als gleich
    return re.sub(r"[^0-9a-zäöüß]", "", text.casefold())


BEREICH_KEYS = {bereich_key(b): b for b in BEREICHE}


def find_file(name: str) -> Path:
    ordner_liste = (ORDNER, BENUTZERGRUPPEN_ORDNER, ORDNER.parent, Path.cwd())
    for ordner in ordner_liste:
        path = ordner / name
        if path.exists():
            return path
    sys.exit(f"Datei nicht gefunden: {name} (gesucht in: {', '.join(map(str, ordner_liste))})")


def read_zuordnung(path: Path):
    """Liefert {gruppe_key: [bereiche]} aus "Benutzergruppen Zuordnung.xlsx"."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        sys.exit(f"{path.name} ist leer.")

    spalten = {}
    for idx, header in enumerate(rows[0]):
        header = clean(header)
        if not header:
            continue
        bereich = BEREICH_KEYS.get(bereich_key(header))
        if bereich:
            spalten[idx] = bereich
        else:
            print(f"Warnung: unbekannte Spalte '{header}' in {path.name} wird ignoriert.")

    gruppen = {}
    for row in rows[1:]:
        for idx, bereich in spalten.items():
            gruppe = clean(row[idx] if idx < len(row) else None)
            if gruppe:
                liste = gruppen.setdefault(gruppe_key(gruppe), [])
                if bereich not in liste:
                    liste.append(bereich)
    return gruppen


def bereiche_fuer_gruppe(gruppe: str, gruppen):
    gefunden = gruppen.get(gruppe_key(gruppe))
    if gefunden:
        return gefunden
    treffer = IT_MUSTER.search(gruppe)
    if treffer:
        ziffer = treffer.group(1)
        bereich = LEITUNG if ziffer == "0" else f"IT {ziffer}"
        if bereich in BEREICHE:
            return [bereich]
    return []


def read_quelle(path: Path, gruppen):
    """Liefert pro Zeile (Zellen A bis H, zugeordnete Bereiche in fester Reihenfolge)."""
    ws = load_workbook(path, data_only=True).active
    zeilen = []
    for cells in ws.iter_rows(min_row=QUELLE_START_ROW, max_col=QUELLE_SPALTEN):
        if all(clean(c.value) == "" for c in cells):
            continue
        gefunden = set()
        for gruppe in clean(cells[COL_GRUPPEN - 1].value).split(","):
            gruppe = clean(gruppe)
            if gruppe:
                gefunden.update(bereiche_fuer_gruppe(gruppe, gruppen))
        bereiche = [b for b in BEREICHE if b in gefunden]
        zeilen.append((cells, bereiche))
    return zeilen


def header_zelle(ws, col: int, text: str) -> None:
    """Schreibt eine Überschrift in Zeile 2 im Stil der Überschrift von Spalte H."""
    vorbild = ws.cell(row=VORLAGE_HEADER_ROW, column=COL_GRUPPEN)
    cell = ws.cell(row=VORLAGE_HEADER_ROW, column=col, value=text)
    if vorbild.has_style:
        cell.font = copy(vorbild.font)
        cell.fill = copy(vorbild.fill)
        cell.border = copy(vorbild.border)
        cell.alignment = copy(vorbild.alignment)
    ws.column_dimensions[cell.column_letter].width = max(len(text) + 4, 20)


def write_liste(vorlage: Path, ausgabe: Path, zeilen, mit_zuordnung: bool) -> None:
    wb = load_workbook(vorlage)
    ws = wb.active

    header_zelle(ws, COL_DOPPELT, "Doppelzuordnung")
    if mit_zuordnung:
        header_zelle(ws, COL_ZUORDNUNG, "Zuordnung")

    for offset, (cells, bereiche) in enumerate(zeilen):
        row = VORLAGE_START_ROW + offset
        for col, quelle in enumerate(cells, start=1):
            ziel = ws.cell(row=row, column=col, value=quelle.value)
            ziel.number_format = quelle.number_format
        if len(bereiche) > 1:
            ws.cell(row=row, column=COL_DOPPELT, value=", ".join(bereiche))
        if mit_zuordnung:
            ws.cell(row=row, column=COL_ZUORDNUNG, value=", ".join(bereiche))

    wb.save(ausgabe)


def main() -> None:
    quelle_datei = find_file(QUELLE_DATEI)
    zuordnung_datei = find_file(ZUORDNUNG_DATEI)
    vorlage_datei = find_file(VORLAGE_DATEI)

    gruppen = read_zuordnung(zuordnung_datei)
    zeilen = read_quelle(quelle_datei, gruppen)

    print(f"{len(zeilen)} Zeilen gelesen.")
    for bereich, kuerzel in BEREICHE.items():
        auswahl = [z for z in zeilen if bereich in z[1]]
        write_liste(vorlage_datei, AUSGABE_ORDNER / AUSGABE_MUSTER.format(kuerzel), auswahl, mit_zuordnung=False)
        print(f"  {kuerzel}: {len(auswahl)}")

    write_liste(vorlage_datei, AUSGABE_ORDNER / AUSGABE_MUSTER.format(GESAMT), zeilen, mit_zuordnung=True)
    print(f"  {GESAMT}: {len(zeilen)}")

    doppelt = sum(1 for _, b in zeilen if len(b) > 1)
    ohne = sum(1 for _, b in zeilen if not b)
    print(f"Doppelzuordnungen: {doppelt}")
    if ohne:
        print(f"Ohne Zuordnung (nur in der Gesamtliste): {ohne}")
    print(f"Fertig: 10 Dateien in {AUSGABE_ORDNER}")
    print(f"Log: {LOG_DATEI}")


class Tee:
    """Schreibt alles gleichzeitig ins Terminal und in die Logdatei."""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, text):
        for stream in self.streams:
            stream.write(text)

    def flush(self):
        for stream in self.streams:
            stream.flush()


def main_mit_log() -> None:
    AUSGABE_ORDNER.mkdir(exist_ok=True)
    with open(LOG_DATEI, "w", encoding="utf-8") as log:
        stdout, stderr = sys.stdout, sys.stderr
        sys.stdout, sys.stderr = Tee(stdout, log), Tee(stderr, log)
        try:
            print(f"Lauf vom {datetime.now():%d.%m.%Y %H:%M:%S}")
            main()
        except SystemExit as exc:
            # sys.exit("Meldung") gibt die Meldung sonst erst nach dem Schließen der Logdatei aus
            if isinstance(exc.code, str):
                print(exc.code, file=sys.stderr)
                raise SystemExit(1)
            raise
        except Exception:
            traceback.print_exc()
            raise SystemExit(1)
        finally:
            sys.stdout, sys.stderr = stdout, stderr


if __name__ == "__main__":
    main_mit_log()
