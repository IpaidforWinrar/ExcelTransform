"""
Vergleicht Spalte A mit Spalte B einer Exceltabelle (ab Zeile 3) und listet die Zuordnungspaare.

Spalte A enthält die Bereiche (IT 1 bis IT 6, "IT-Leitung + Stab", "übergreifend"),
Spalte B die zugeordneten Begriffe.

Ausgabe ist eine neue Exceltabelle mit einer Spalte pro Bereich, darunter alle Begriffe,
die diesem Bereich zugeordnet sind. Die Extraspalte "doppelt" listet jeden Begriff,
der mehr als einmal vorkommt, zusammen mit den Bereichen, denen er zugeordnet ist,
z. B. "Gruppe X → IT 1, IT 3" oder "Gruppe Y → IT 2 (2x)".

Aufruf:
    python zuordnung.py [Eingabedatei] [Ausgabedatei]

Ohne Eingabedatei wird "test.xlsx" im selben Ordner wie dieses Skript verwendet.
Ohne Ausgabedatei wird "<Eingabedatei>_Zuordnung.xlsx" gespeichert.
"""

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

DEFAULT_INPUT = Path(__file__).resolve().parent / "test.xlsx"

START_ROW = 3

BEREICHE = [
    "IT 1",
    "IT 2",
    "IT 3",
    "IT 4",
    "IT 5",
    "IT 6",
    "IT-Leitung + Stab",
    "übergreifend",
]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_key(text: str) -> str:
    # "IT1", "it 1" und "IT  1" werden alle als "it 1" erkannt
    return re.sub(r"^it\s*(\d)$", r"it \1", text.lower())


BEREICH_KEYS = {normalize_key(b): b for b in BEREICHE}


def read_pairs(input_path: Path):
    ws = load_workbook(input_path, read_only=True, data_only=True).active
    pairs = []
    skipped = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=START_ROW, max_col=2, values_only=True), start=START_ROW
    ):
        a = clean(row[0] if len(row) > 0 else None)
        b = clean(row[1] if len(row) > 1 else None)
        if not a and not b:
            continue
        if not a or not b:
            skipped.append(row_idx)
            continue
        bereich = BEREICH_KEYS.get(normalize_key(a), a)
        pairs.append((bereich, b))
    return pairs, skipped


def build_output(pairs, output_path: Path):
    # Begriffe pro Bereich, ohne Wiederholungen, in der Reihenfolge des Auftretens
    per_bereich = defaultdict(list)
    # Wie oft kommt ein Begriff in welchem Bereich vor?
    vorkommen = defaultdict(Counter)

    for bereich, begriff in pairs:
        if begriff not in per_bereich[bereich]:
            per_bereich[bereich].append(begriff)
        vorkommen[begriff][bereich] += 1

    # Unbekannte Bereiche gehen nicht verloren, sondern bekommen eigene Spalten
    unbekannt = [b for b in per_bereich if b not in BEREICHE]
    spalten = BEREICHE + unbekannt

    doppelt = []
    for begriff, counter in vorkommen.items():
        if sum(counter.values()) < 2:
            continue
        teile = [
            f"{bereich} ({anzahl}x)" if anzahl > 1 else bereich
            for bereich, anzahl in sorted(counter.items(), key=lambda kv: spalten.index(kv[0]))
        ]
        doppelt.append(f"{begriff} → {', '.join(teile)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Zuordnung"

    headers = spalten + ["doppelt"]
    columns = [per_bereich[b] for b in spalten] + [doppelt]

    for col_idx, (header, values) in enumerate(zip(headers, columns), start=1):
        ws.cell(row=1, column=col_idx, value=header).font = Font(bold=True)
        for row_idx, value in enumerate(values, start=2):
            ws.cell(row=row_idx, column=col_idx, value=value)
        width = max([len(header)] + [len(v) for v in values]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)

    ws.freeze_panes = "A2"
    wb.save(output_path)
    return unbekannt, doppelt


def main() -> None:
    input_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT
    if len(sys.argv) > 2:
        output_path = Path(sys.argv[2])
    else:
        output_path = input_path.with_name(f"{input_path.stem}_Zuordnung.xlsx")

    if not input_path.exists():
        sys.exit(f"Datei nicht gefunden: {input_path}")

    pairs, skipped = read_pairs(input_path)
    unbekannt, doppelt = build_output(pairs, output_path)

    print(f"{len(pairs)} Zuordnungspaare gefunden, {len(doppelt)} Begriffe doppelt.")
    if unbekannt:
        print(f"Warnung: unbekannte Bereiche in Spalte A (als eigene Spalten ausgegeben): {', '.join(unbekannt)}")
    if skipped:
        print(f"Warnung: Zeilen mit nur einem Wert in A oder B übersprungen: {', '.join(map(str, skipped))}")
    print(f"Fertig: {output_path}")


if __name__ == "__main__":
    main()
