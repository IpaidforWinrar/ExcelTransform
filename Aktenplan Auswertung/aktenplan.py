"""
Findet Zuschlüsselungen im Aktenplan: Welche Einträge aus Spalte C sind genau einem Bereich
zugeordnet und welche mehreren?

Eingabe: "AktenplanAlt.xlsx", Werte ab Zeile 3.
    Spalte B: Bereich (IT 1 bis IT 7, IT-Leitung + Stab, übergreifend)
    Spalte C: Aktenplan-Eintrag, z. B. "IT.01.01.01 Inspire"

Ausgabe: "Aktenplan Zuordnung.xlsx" mit einer Zeile pro Eintrag aus Spalte C:
    Aktenplan-Eintrag | Status | Bereiche | Anzahl Bereiche | Anzahl Zeilen | Zeilen
Status ist "unique", wenn der Eintrag nur einem Bereich zugeordnet ist, sonst "mehrfach".
Mehrfach zugeordnete Einträge stehen oben. Bei "Bereiche" steht dahinter, wie oft der
Eintrag dem Bereich zugeordnet ist, z. B. "IT 1 (2x), IT 7".

Einträge in Spalte C werden ohne Rücksicht auf Groß-/Kleinschreibung und überzählige
Leerzeichen verglichen.

Aufruf:
    python aktenplan.py

"AktenplanAlt.xlsx" wird im Ordner dieses Skripts gesucht, danach im aktuellen Ordner.
Das Ergebnis landet im selben Ordner wie "AktenplanAlt.xlsx".
"""

import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ORDNER = Path(__file__).resolve().parent

QUELLE_DATEI = "AktenplanAlt.xlsx"
AUSGABE_NAME = "Aktenplan Zuordnung.xlsx"

START_ROW = 3
COL_BEREICH = 2  # Spalte B
COL_EINTRAG = 3  # Spalte C

BEREICHE = [
    "IT 1",
    "IT 2",
    "IT 3",
    "IT 4",
    "IT 5",
    "IT 6",
    "IT 7",
    "IT-Leitung + Stab",
    "übergreifend",
]

STATUS_UNIQUE = "unique"
STATUS_MEHRFACH = "mehrfach"
STATUS_FARBEN = {STATUS_UNIQUE: "C6EFCE", STATUS_MEHRFACH: "FFC7CE"}

HEADERS = ["Aktenplan-Eintrag", "Status", "Bereiche", "Anzahl Bereiche", "Anzahl Zeilen", "Zeilen"]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def bereich_key(text: str) -> str:
    # "IT 1" und "IT1" sowie "IT-Leitung + Stab" und "IT Leitung + Stab" gelten als gleich
    return re.sub(r"[^0-9a-zäöüß]", "", text.casefold())


BEREICH_KEYS = {bereich_key(b): b for b in BEREICHE}


def bereich_name(text: str) -> str:
    return BEREICH_KEYS.get(bereich_key(text), text)


def sort_key_bereich(bereich: str):
    return (BEREICHE.index(bereich) if bereich in BEREICHE else len(BEREICHE), bereich)


def find_file(name: str) -> Path:
    for ordner in (ORDNER, Path.cwd()):
        path = ordner / name
        if path.exists():
            return path
    sys.exit(f"Datei nicht gefunden: {name} (gesucht in {ORDNER} und {Path.cwd()})")


def read_eintraege(path: Path):
    """Liefert {eintrag_key: {"eintrag", "bereiche": Counter, "zeilen": [...]}} und übersprungene Zeilen."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    eintraege = {}
    ohne_bereich = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=START_ROW, max_col=COL_EINTRAG, values_only=True), start=START_ROW
    ):
        bereich = clean(row[COL_BEREICH - 1] if len(row) >= COL_BEREICH else None)
        eintrag = clean(row[COL_EINTRAG - 1] if len(row) >= COL_EINTRAG else None)
        if not eintrag:
            continue
        if not bereich:
            ohne_bereich.append(row_idx)
            continue
        e = eintraege.setdefault(
            eintrag.casefold(), {"eintrag": eintrag, "bereiche": Counter(), "zeilen": []}
        )
        e["bereiche"][bereich_name(bereich)] += 1
        e["zeilen"].append(row_idx)
    return eintraege, ohne_bereich


def build_rows(eintraege):
    rows = []
    for e in eintraege.values():
        bereiche = sorted(e["bereiche"], key=sort_key_bereich)
        bereiche_text = ", ".join(
            f"{b} ({e['bereiche'][b]}x)" if e["bereiche"][b] > 1 else b for b in bereiche
        )
        status = STATUS_UNIQUE if len(bereiche) == 1 else STATUS_MEHRFACH
        rows.append((
            e["eintrag"],
            status,
            bereiche_text,
            len(bereiche),
            len(e["zeilen"]),
            ", ".join(map(str, e["zeilen"])),
        ))
    rows.sort(key=lambda r: (r[1] != STATUS_MEHRFACH, r[0].casefold()))
    return rows


def write_output(rows, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Aktenplan Zuordnung"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)
        ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", start_color=STATUS_FARBEN[row[1]])

    for col_idx, header in enumerate(HEADERS, start=1):
        width = max([len(header)] + [len(str(r[col_idx - 1])) for r in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    quelle_datei = find_file(QUELLE_DATEI)
    ausgabe_datei = quelle_datei.with_name(AUSGABE_NAME)

    eintraege, ohne_bereich = read_eintraege(quelle_datei)
    rows = build_rows(eintraege)
    write_output(rows, ausgabe_datei)

    mehrfach = [r for r in rows if r[1] == STATUS_MEHRFACH]
    print(f"{len(rows)} verschiedene Einträge in Spalte C gefunden.")
    print(f"  {STATUS_UNIQUE}: {len(rows) - len(mehrfach)}")
    print(f"  {STATUS_MEHRFACH}: {len(mehrfach)}")
    for r in mehrfach:
        print(f"    {r[0]} -> {r[2]}")
    if ohne_bereich:
        print(f"Warnung: Zeilen ohne Bereich in Spalte B übersprungen: {', '.join(map(str, ohne_bereich))}")
    print(f"Fertig: {ausgabe_datei}")


if __name__ == "__main__":
    main()
