"""
Vergleicht "filled Auswertung Benutzergruppen ZB IT.xlsx" (dieses Jahr) mit
"toCompareLastYear.xlsx" (letztes Jahr) und listet die Neuzugänge.

Beide Tabellen haben denselben Aufbau: A = betrifft, B = Gruppe, C = Name, Werte ab Zeile 3.

Ein Neuzugang ist eine Kombination aus Gruppe und Name, die dieses Jahr vorkommt,
letztes Jahr aber nicht. Groß-/Kleinschreibung und überzählige Leerzeichen werden ignoriert.
Ändert sich nur "betrifft", zählt das nicht als Neuzugang.

Ergebnis ist "Neuzugänge.xlsx" mit den Spalten betrifft, Gruppe, Name.

Aufruf:
    python comparator.py

Die Exceldateien werden im Ordner dieses Skripts gesucht, danach im aktuellen Ordner.
Das Ergebnis landet im selben Ordner wie die Datei von diesem Jahr.
"""

import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ORDNER = Path(__file__).resolve().parent

AKTUELL_DATEI = "filled Auswertung Benutzergruppen ZB IT.xlsx"
VORJAHR_DATEI = "toCompareLastYear.xlsx"
AUSGABE_NAME = "Neuzugänge.xlsx"

START_ROW = 3
HEADERS = ["betrifft", "Gruppe", "Name"]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def find_file(name: str) -> Path:
    for ordner in (ORDNER, Path.cwd()):
        path = ordner / name
        if path.exists():
            return path
    sys.exit(f"Datei nicht gefunden: {name} (gesucht in {ORDNER} und {Path.cwd()})")


def read_rows(path: Path):
    """Liefert (betrifft, gruppe, name) für jede Zeile ab START_ROW mit Gruppe oder Name."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = []
    for row in ws.iter_rows(min_row=START_ROW, max_col=3, values_only=True):
        betrifft, gruppe, name = (clean(row[i]) if i < len(row) else "" for i in range(3))
        if gruppe or name:
            rows.append((betrifft, gruppe, name))
    return rows


def key(gruppe: str, name: str):
    return (gruppe.casefold(), name.casefold())


def main() -> None:
    aktuell_datei = find_file(AKTUELL_DATEI)
    vorjahr_datei = find_file(VORJAHR_DATEI)
    ausgabe_datei = aktuell_datei.with_name(AUSGABE_NAME)

    aktuell = read_rows(aktuell_datei)
    vorjahr = read_rows(vorjahr_datei)
    vorjahr_keys = {key(gruppe, name) for _, gruppe, name in vorjahr}

    neuzugaenge = []
    gesehen = set()
    for betrifft, gruppe, name in aktuell:
        k = key(gruppe, name)
        if k in vorjahr_keys or k in gesehen:
            continue
        gesehen.add(k)
        neuzugaenge.append((betrifft, gruppe, name))

    wb = Workbook()
    ws = wb.active
    ws.title = "Neuzugänge"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in neuzugaenge:
        ws.append(row)

    for col_idx, header in enumerate(HEADERS, start=1):
        width = max([len(header)] + [len(r[col_idx - 1]) for r in neuzugaenge]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ausgabe_datei)

    print(f"Dieses Jahr: {len(aktuell)} Zeilen, letztes Jahr: {len(vorjahr)} Zeilen")
    print(f"Neuzugänge: {len(neuzugaenge)}")
    print(f"Fertig: {ausgabe_datei}")


if __name__ == "__main__":
    main()
