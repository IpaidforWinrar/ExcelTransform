"""
Gleicht test_Zuordnung.xlsx (Ergebnis von zuordnung.py) mit matching.xlsx ab.

test_Zuordnung.xlsx: Zeile 1 enthält die Bereiche (IT 1 bis IT 6 usw.),
darunter stehen die Begriffe, die dem jeweiligen Bereich zugeordnet sind.

matching.xlsx: ab Zeile 2 steht in Spalte A der Begriff und in Spalte F der Bereich.

Verglichen werden die Paare (Bereich, Begriff). Groß-/Kleinschreibung und
überzählige Leerzeichen werden dabei ignoriert. Ausgabe ist Abgleich.xlsx mit
einer Zeile pro Paar und dem Status:
    - "in beiden"
    - "nur in test_Zuordnung"
    - "nur in matching"
Die Spalte "In der anderen Liste bei" zeigt, welchem Bereich der Begriff in der
anderen Liste zugeordnet ist, falls das Paar dort fehlt.

Aufruf:
    python abgleich.py

Alle Dateien liegen im selben Ordner wie dieses Skript.
"""

import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from zuordnung import BEREICH_KEYS, BEREICHE, clean, normalize_key

ORDNER = Path(__file__).resolve().parent
ZUORDNUNG_DATEI = ORDNER / "test_Zuordnung.xlsx"
MATCHING_DATEI = ORDNER / "matching.xlsx"
AUSGABE_DATEI = ORDNER / "Abgleich.xlsx"

MATCHING_START_ROW = 2
MATCHING_COL_BEGRIFF = 1  # Spalte A
MATCHING_COL_BEREICH = 6  # Spalte F

STATUS_BEIDE = "in beiden"
STATUS_NUR_ZUORDNUNG = "nur in test_Zuordnung"
STATUS_NUR_MATCHING = "nur in matching"

STATUS_FARBEN = {
    STATUS_BEIDE: "C6EFCE",
    STATUS_NUR_ZUORDNUNG: "FFEB9C",
    STATUS_NUR_MATCHING: "FFC7CE",
}


def bereich_name(text: str) -> str:
    return BEREICH_KEYS.get(normalize_key(text), text)


def begriff_key(text: str) -> str:
    return text.casefold()


def read_zuordnung(path: Path):
    """Liefert {(bereich, begriff_key): begriff} aus test_Zuordnung.xlsx."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    pairs = {}
    if not rows:
        return pairs
    headers = [clean(h) for h in rows[0]]
    for row in rows[1:]:
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            begriff = clean(row[col_idx] if col_idx < len(row) else None)
            if begriff:
                pairs[(bereich_name(header), begriff_key(begriff))] = begriff
    return pairs


def read_matching(path: Path):
    """Liefert {(bereich, begriff_key): begriff} aus matching.xlsx und übersprungene Zeilen."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    pairs = {}
    skipped = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=MATCHING_START_ROW, max_col=MATCHING_COL_BEREICH, values_only=True),
        start=MATCHING_START_ROW,
    ):
        begriff = clean(row[MATCHING_COL_BEGRIFF - 1] if len(row) >= MATCHING_COL_BEGRIFF else None)
        bereich = clean(row[MATCHING_COL_BEREICH - 1] if len(row) >= MATCHING_COL_BEREICH else None)
        if not begriff and not bereich:
            continue
        if not begriff or not bereich:
            skipped.append(row_idx)
            continue
        pairs[(bereich_name(bereich), begriff_key(begriff))] = begriff
    return pairs, skipped


def bereiche_pro_begriff(pairs):
    result = defaultdict(set)
    for bereich, key in pairs:
        result[key].add(bereich)
    return result


def sort_key_bereich(bereich: str):
    return (BEREICHE.index(bereich) if bereich in BEREICHE else len(BEREICHE), bereich)


def compare(zuordnung, matching):
    bereiche_zuordnung = bereiche_pro_begriff(zuordnung)
    bereiche_matching = bereiche_pro_begriff(matching)

    rows = []
    for pair in set(zuordnung) | set(matching):
        bereich, key = pair
        if pair in zuordnung and pair in matching:
            status, andere = STATUS_BEIDE, set()
        elif pair in zuordnung:
            status, andere = STATUS_NUR_ZUORDNUNG, bereiche_matching.get(key, set())
        else:
            status, andere = STATUS_NUR_MATCHING, bereiche_zuordnung.get(key, set())
        begriff = zuordnung.get(pair) or matching.get(pair)
        andere_text = ", ".join(sorted(andere, key=sort_key_bereich))
        rows.append((begriff, bereich, status, andere_text))

    status_order = [STATUS_NUR_ZUORDNUNG, STATUS_NUR_MATCHING, STATUS_BEIDE]
    rows.sort(key=lambda r: (status_order.index(r[2]), sort_key_bereich(r[1]), r[0].casefold()))
    return rows


def write_output(rows, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Abgleich"

    headers = ["Begriff", "Bereich", "Status", "In der anderen Liste bei"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)
        ws.cell(row=ws.max_row, column=3).fill = PatternFill(
            "solid", start_color=STATUS_FARBEN[row[2]]
        )

    for col_idx, header in enumerate(headers, start=1):
        width = max([len(header)] + [len(str(r[col_idx - 1])) for r in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width, 80)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    for path in (ZUORDNUNG_DATEI, MATCHING_DATEI):
        if not path.exists():
            sys.exit(f"Datei nicht gefunden: {path}")

    zuordnung = read_zuordnung(ZUORDNUNG_DATEI)
    matching, skipped = read_matching(MATCHING_DATEI)
    rows = compare(zuordnung, matching)
    write_output(rows, AUSGABE_DATEI)

    for status in (STATUS_BEIDE, STATUS_NUR_ZUORDNUNG, STATUS_NUR_MATCHING):
        print(f"{status}: {sum(1 for r in rows if r[2] == status)}")
    if skipped:
        print(f"Warnung: Zeilen in matching.xlsx mit nur einem Wert in A oder F übersprungen: {', '.join(map(str, skipped))}")
    print(f"Fertig: {AUSGABE_DATEI}")


if __name__ == "__main__":
    main()
